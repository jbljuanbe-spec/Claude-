#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Motor de render de los CVs.

El contenido se declara una sola vez llamando a section()/role()/bullet()/... sobre
un objeto CVDoc, que sólo recoge bloques. Después cada bloque se pinta dos veces:
en DOCX (python-docx, para portales ATS y para editar en Word) y en PDF (HTML
renderizado con Chromium headless, para enviar por email).
"""

import html
import os
import subprocess

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
os.makedirs(OUT, exist_ok=True)

CHROMIUM = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

ACCENT = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x44, 0x44, 0x44)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT_HEX = "#1F3864"
GREY_HEX = "#444444"
FONT = "Calibri"
RIGHT_TAB = Cm(17.4)


# ============================================================ recogida de bloques

class CVDoc:
    def __init__(self):
        self.blocks = []

    def add(self, kind, **kw):
        self.blocks.append((kind, kw))

    def save(self, docx_path):
        render_docx(self, docx_path)
        render_pdf(self, docx_path.replace(".docx", ".pdf"))


def new_doc():
    return CVDoc()


def header_block(doc, name, tagline, contact):
    doc.add("header", name=name, tagline=tagline, contact=contact)


def section(doc, title):
    doc.add("section", title=title)


def body(doc, text, size=9.8, italic=False, space_after=2, color=BLACK):
    doc.add("body", text=text, size=size, italic=italic,
            note=(color is GREY), space_after=space_after)


def role(doc, title, dates, org):
    doc.add("role", title=title, dates=dates, org=org)


def bullet(doc, text, bold_lead=None):
    doc.add("bullet", text=text, lead=bold_lead)


def edu(doc, title, dates, org, detail=None):
    doc.add("edu", title=title, dates=dates, org=org, detail=detail)


def skill_line(doc, label, text):
    doc.add("skill", label=label, text=text)


# ============================================================ render DOCX

def set_font(run, size=10, bold=False, italic=False, color=BLACK, caps=False, spacing=None):
    """El orden de los hijos de w:rPr lo fija el esquema OOXML; por eso todo se
    inserta vía python-docx salvo w:spacing, que se coloca a mano antes de w:sz."""
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True
    if spacing:
        rpr = run._element.get_or_add_rPr()
        s = OxmlElement('w:spacing')
        s.set(qn('w:val'), str(int(spacing * 20)))
        sz = rpr.find(qn('w:sz'))
        if sz is not None:
            sz.addprevious(s)
        else:
            rpr.append(s)


def _para(d, space_before=0, space_after=2, line=1.06, left=0, hanging=None):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.widow_control = True
    if left:
        pf.left_indent = Cm(left)
    if hanging:
        pf.first_line_indent = Cm(-hanging)
    return p


def _bottom_border(p, color="C6CEDF", size=6):
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(size))
    b.set(qn('w:space'), '2')
    b.set(qn('w:color'), color)
    pbdr.append(b)
    pPr.append(pbdr)


def render_docx(cv, path):
    d = Document()
    s = d.sections[0]
    s.top_margin, s.bottom_margin = Cm(1.3), Cm(1.3)
    s.left_margin, s.right_margin = Cm(1.8), Cm(1.8)
    st = d.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(9.8)
    st.paragraph_format.space_after = Pt(2)

    for kind, k in cv.blocks:
        if kind == "header":
            p = _para(d, space_after=1)
            set_font(p.add_run(k["name"]), size=20, bold=True, color=ACCENT, spacing=0.8)
            p = _para(d, space_after=3)
            set_font(p.add_run(k["tagline"]), size=10.5, color=GREY)
            lines = k["contact"].split("\n")
            for i, ln in enumerate(lines):
                last = i == len(lines) - 1
                p = _para(d, space_after=8 if last else 1)
                set_font(p.add_run(ln), size=9, color=GREY)
            _bottom_border(p)

        elif kind == "section":
            p = _para(d, space_before=9, space_after=4)
            set_font(p.add_run(k["title"]), size=10, bold=True, color=ACCENT, caps=True, spacing=1.1)
            _bottom_border(p)

        elif kind == "body":
            p = _para(d, space_after=k["space_after"], line=1.1)
            set_font(p.add_run(k["text"]), size=k["size"], italic=k["italic"],
                     color=GREY if k["note"] else BLACK)

        elif kind == "role":
            p = _para(d, space_before=5, space_after=0)
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            set_font(p.add_run(k["title"]), size=10.5, bold=True, color=BLACK)
            set_font(p.add_run("\t"), size=9, color=GREY)
            set_font(p.add_run(k["dates"]), size=9, bold=True, color=GREY)
            p = _para(d, space_after=2)
            set_font(p.add_run(k["org"]), size=9.5, italic=True, color=ACCENT)

        elif kind == "bullet":
            p = _para(d, space_after=1.5, line=1.1, left=0.45, hanging=0.45)
            set_font(p.add_run("▪  "), size=9.8, color=ACCENT)
            if k["lead"]:
                set_font(p.add_run(k["lead"]), size=9.8, bold=True, color=BLACK)
            set_font(p.add_run(k["text"]), size=9.8, color=BLACK)

        elif kind == "edu":
            p = _para(d, space_before=4, space_after=0)
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            set_font(p.add_run(k["title"]), size=9.8, bold=True, color=BLACK)
            set_font(p.add_run("\t"), size=9, color=GREY)
            set_font(p.add_run(k["dates"]), size=9, color=GREY)
            p = _para(d, space_after=1)
            set_font(p.add_run(k["org"]), size=9.3, italic=True, color=ACCENT)
            if k["detail"]:
                p = _para(d, space_after=1, line=1.1)
                set_font(p.add_run(k["detail"]), size=9.3, color=GREY)

        elif kind == "skill":
            p = _para(d, space_after=2, line=1.1, left=0.45, hanging=0.45)
            set_font(p.add_run("▪  "), size=9.8, color=ACCENT)
            set_font(p.add_run(k["label"] + ": "), size=9.8, bold=True, color=BLACK)
            set_font(p.add_run(k["text"]), size=9.8, color=BLACK)

    d.save(path)


# ============================================================ render PDF (HTML + Chromium)

CSS = """
@page { size: A4; margin: 13mm 18mm; }
* { box-sizing: border-box; }
body { font-family: "Liberation Sans", Arial, Helvetica, sans-serif;
       font-size: 9.8pt; line-height: 1.28; color: #1A1A1A; margin: 0; }
.name { font-size: 20pt; font-weight: bold; color: %(accent)s; letter-spacing: .8pt;
        margin: 0 0 1pt; }
.tagline { font-size: 10.5pt; color: %(grey)s; margin: 0 0 3pt; }
.contact { font-size: 9pt; color: %(grey)s; margin: 0 0 8pt; padding-bottom: 3pt;
           border-bottom: .75pt solid #C6CEDF; }
h2 { font-size: 10pt; font-weight: bold; color: %(accent)s; text-transform: uppercase;
     letter-spacing: 1.1pt; margin: 9pt 0 4pt; padding-bottom: 2pt;
     border-bottom: .75pt solid #C6CEDF; }
p { margin: 0 0 2pt; }
.note { font-size: 9.3pt; color: %(grey)s; }
.role { display: flex; justify-content: space-between; align-items: baseline;
        gap: 10pt; margin: 5pt 0 0; break-inside: avoid; }
.role .t { font-size: 10.5pt; font-weight: bold; }
.role .d { font-size: 9pt; font-weight: bold; color: %(grey)s; white-space: nowrap; }
.org { font-size: 9.5pt; font-style: italic; color: %(accent)s; margin: 0 0 2pt; }
.edu { display: flex; justify-content: space-between; align-items: baseline;
       gap: 10pt; margin: 4pt 0 0; break-inside: avoid; }
.edu .t { font-size: 9.8pt; font-weight: bold; }
.edu .d { font-size: 9pt; color: %(grey)s; white-space: nowrap; }
.eorg { font-size: 9.3pt; font-style: italic; color: %(accent)s; margin: 0 0 1pt; }
li { list-style: none; position: relative; padding-left: 12pt; margin: 0 0 1.5pt; }
li:before { content: "\\25AA"; position: absolute; left: 0; color: %(accent)s; }
ul { margin: 0; padding: 0; }
h2, .role, .edu { break-after: avoid; }
li { break-inside: avoid; }
""" % {"accent": ACCENT_HEX, "grey": GREY_HEX}


def e(t):
    return html.escape(t)


def render_pdf(cv, path):
    out, in_list = [], False

    def close():
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    def open_list():
        nonlocal in_list
        if not in_list:
            out.append("<ul>")
            in_list = True

    for kind, k in cv.blocks:
        if kind in ("header", "section", "body", "role", "edu"):
            close()
        if kind == "header":
            out.append('<p class="name">%s</p><p class="tagline">%s</p>'
                       '<p class="contact">%s</p>'
                       % (e(k["name"]), e(k["tagline"]), e(k["contact"]).replace("\n", "<br>")))
        elif kind == "section":
            out.append("<h2>%s</h2>" % e(k["title"]))
        elif kind == "body":
            out.append('<p%s>%s</p>' % (' class="note"' if k["note"] else "", e(k["text"])))
        elif kind == "role":
            out.append('<div class="role"><span class="t">%s</span><span class="d">%s</span></div>'
                       '<p class="org">%s</p>' % (e(k["title"]), e(k["dates"]), e(k["org"])))
        elif kind == "edu":
            out.append('<div class="edu"><span class="t">%s</span><span class="d">%s</span></div>'
                       '<p class="eorg">%s</p>' % (e(k["title"]), e(k["dates"]), e(k["org"])))
            if k["detail"]:
                out.append('<p class="note">%s</p>' % e(k["detail"]))
        elif kind == "bullet":
            open_list()
            lead = "<b>%s</b>" % e(k["lead"]) if k["lead"] else ""
            out.append("<li>%s%s</li>" % (lead, e(k["text"])))
        elif kind == "skill":
            open_list()
            out.append("<li><b>%s: </b>%s</li>" % (e(k["label"]), e(k["text"])))
    close()

    doc = ("<!doctype html><html lang='es'><head><meta charset='utf-8'>"
           "<style>%s</style></head><body>%s</body></html>" % (CSS, "".join(out)))
    tmp = path.replace(".pdf", ".html")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(doc)
    subprocess.run(
        [CHROMIUM, "--headless", "--disable-gpu", "--no-sandbox", "--no-pdf-header-footer",
         "--print-to-pdf=" + path, "file://" + tmp],
        check=True, capture_output=True, timeout=120)
    if not os.environ.get("KEEP_HTML"): os.remove(tmp)
